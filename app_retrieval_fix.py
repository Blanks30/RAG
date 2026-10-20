
import os
import io
import csv
import json
import xml.etree.ElementTree as ET
import urllib.request
import urllib.error

import chromadb
import pymupdf
import pytesseract
import streamlit as st

from PIL import Image
from bs4 import BeautifulSoup
from dotenv import load_dotenv
from sentence_transformers import SentenceTransformer
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook


load_dotenv(override=True)

st.set_page_config(
    page_title="Multi-Format RAG Chatbot",
    page_icon="📚",
    layout="centered",
)

st.title("Multi-Format RAG Chatbot")
st.write(
    "Upload documents, spreadsheets, presentations, text files, "
    "HTML, or images and ask questions about their contents."
)

OLLAMA_URL = os.getenv(
    "OLLAMA_BASE_URL", "http://localhost:11434"
).rstrip("/")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "qwen2.5:3b")
OLLAMA_KEEP_ALIVE = os.getenv("OLLAMA_KEEP_ALIVE", "30m")

TESSERACT_PATH = os.getenv(
    "TESSERACT_CMD",
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
)
pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH

SUPPORTED_TYPES = [
    "pdf", "docx", "pptx", "xlsx",
    "html", "htm", "txt", "md",
    "csv", "json", "xml",
    "png", "jpg", "jpeg", "bmp", "tif", "tiff", "webp",
    "py", "js", "ts", "jsx", "tsx", "java",
    "c", "cpp", "h", "hpp", "cs", "go", "rs",
    "sql", "css", "yaml", "yml", "toml", "ini",
    "log", "sh", "bat", "ps1",
]

IMAGE_TYPES = {
    "png", "jpg", "jpeg", "bmp", "tif", "tiff", "webp"
}

TEXT_TYPES = {
    "txt", "md", "csv", "json", "xml",
    "py", "js", "ts", "jsx", "tsx", "java",
    "c", "cpp", "h", "hpp", "cs", "go", "rs",
    "sql", "css", "yaml", "yml", "toml", "ini",
    "log", "sh", "bat", "ps1",
}


@st.cache_resource
def get_model():
    return SentenceTransformer("all-MiniLM-L6-v2")


@st.cache_resource
def get_db():
    return chromadb.PersistentClient(path="./chroma_db")


def extract_image_text(image_bytes):
    """Extract readable text from an image using Tesseract OCR."""
    with Image.open(io.BytesIO(image_bytes)) as image:
        image = image.convert("RGB")
        return pytesseract.image_to_string(
            image, lang="eng"
        ).strip()


def extract_pdf(file_bytes, filename):
    """Extract PDF text and use OCR on pages without readable text."""
    pdf = pymupdf.open(stream=file_bytes, filetype="pdf")
    records = []

    try:
        for page_number, page in enumerate(pdf, start=1):
            text = " ".join(page.get_text().split())

            if not text:
                pixmap = page.get_pixmap(
                    matrix=pymupdf.Matrix(2, 2),
                    alpha=False,
                )
                image_bytes = pixmap.tobytes("png")
                text = extract_image_text(image_bytes)
                location = f"Page {page_number} (OCR)"
            else:
                location = f"Page {page_number}"

            if text:
                records.append({
                    "text": text,
                    "source": filename,
                    "location": location,
                })
    finally:
        pdf.close()

    return records


def extract_document(file_bytes, filename):
    """Extract text and source locations from a supported file."""
    extension = filename.rsplit(".", 1)[-1].lower()
    records = []

    if extension == "pdf":
        return extract_pdf(file_bytes, filename)

    if extension in IMAGE_TYPES:
        text = extract_image_text(file_bytes)
        if text:
            records.append({
                "text": text,
                "source": filename,
                "location": "Image OCR",
            })
        return records

    if extension in {"html", "htm"}:
        soup = BeautifulSoup(file_bytes, "html.parser")
        for element in soup([
            "script", "style", "noscript", "nav", "footer"
        ]):
            element.decompose()

        text = " ".join(soup.get_text(" ").split())
        if text:
            records.append({
                "text": text,
                "source": filename,
                "location": "HTML content",
            })
        return records

    if extension == "docx":
        document = Document(io.BytesIO(file_bytes))

        for index, paragraph in enumerate(
            document.paragraphs, start=1
        ):
            text = paragraph.text.strip()
            if text:
                records.append({
                    "text": text,
                    "source": filename,
                    "location": f"Paragraph {index}",
                })

        for table_index, table in enumerate(
            document.tables, start=1
        ):
            for row_index, row in enumerate(table.rows, start=1):
                text = " | ".join(
                    cell.text.strip() for cell in row.cells
                    if cell.text.strip()
                )
                if text:
                    records.append({
                        "text": text,
                        "source": filename,
                        "location": (
                            f"Table {table_index}, row {row_index}"
                        ),
                    })
        return records

    if extension == "pptx":
        presentation = Presentation(io.BytesIO(file_bytes))

        for slide_number, slide in enumerate(
            presentation.slides, start=1
        ):
            parts = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    if shape.text.strip():
                        parts.append(shape.text.strip())

                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        parts.append(" | ".join(
                            cell.text.strip()
                            for cell in row.cells
                        ))

            text = "\n".join(parts).strip()
            if text:
                records.append({
                    "text": text,
                    "source": filename,
                    "location": f"Slide {slide_number}",
                })
        return records

    if extension == "xlsx":
        workbook = load_workbook(
            io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
        try:
            for sheet in workbook.worksheets:
                for row_number, row in enumerate(
                    sheet.iter_rows(values_only=True), start=1
                ):
                    values = [
                        str(value).strip()
                        for value in row
                        if value is not None
                        and str(value).strip()
                    ]
                    if values:
                        records.append({
                            "text": " | ".join(values),
                            "source": filename,
                            "location": (
                                f"Sheet {sheet.title}, row {row_number}"
                            ),
                        })
        finally:
            workbook.close()
        return records

    if extension == "csv":
        decoded = file_bytes.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(decoded))
        for row_number, row in enumerate(reader, start=1):
            values = [value.strip() for value in row if value.strip()]
            if values:
                records.append({
                    "text": " | ".join(values),
                    "source": filename,
                    "location": f"Row {row_number}",
                })
        return records

    if extension == "json":
        data = json.loads(file_bytes.decode(
            "utf-8-sig", errors="replace"
        ))
        text = json.dumps(data, ensure_ascii=False, indent=2)
        records.append({
            "text": text,
            "source": filename,
            "location": "JSON content",
        })
        return records

    if extension == "xml":
        root = ET.fromstring(
            file_bytes.decode("utf-8-sig", errors="replace")
        )
        text = " ".join(part.strip() for part in root.itertext()
                        if part.strip())
        if text:
            records.append({
                "text": text,
                "source": filename,
                "location": "XML content",
            })
        return records

    if extension in TEXT_TYPES:
        text = file_bytes.decode("utf-8-sig", errors="replace")
        for line_number, line in enumerate(text.splitlines(), start=1):
            if line.strip():
                records.append({
                    "text": line.strip(),
                    "source": filename,
                    "location": f"Line {line_number}",
                })
        return records

    raise ValueError(f"Unsupported file type: .{extension}")


def split_records(records, chunk_size=1000, overlap=150):
    """Split extracted text into overlapping chunks."""
    chunks = []
    step = chunk_size - overlap

    for record in records:
        text = record["text"].strip()

        for start in range(0, len(text), step):
            chunk = text[start:start + chunk_size].strip()
            if chunk:
                chunks.append({
                    "text": chunk,
                    "source": record["source"],
                    "location": record["location"],
                })

    return chunks


def resolve_requested_source(question, indexed_files):
    """Return an indexed filename when the question names a specific file/type."""
    question_lower = question.lower()

    # Prefer an exact filename if the user included one in the question.
    for filename in indexed_files:
        if filename.lower() in question_lower:
            return filename

    # Resolve common document-type references to supported file extensions.
    type_rules = [
        (("powerpoint", "presentation", "slide deck", "pptx"), (".pptx",)),
        (("excel", "spreadsheet", "workbook", "xlsx"), (".xlsx",)),
        (("word document", "word file", "docx"), (".docx",)),
        (("csv file", "csv"), (".csv",)),
        (("markdown", "markdown file", ".md"), (".md",)),
        (("text file", "txt file", ".txt"), (".txt",)),
        (("html file", "webpage", "html"), (".html", ".htm")),
        (("scanned pdf", "pdf file", "pdf document", "pdf"), (".pdf",)),
        (("json file", "json"), (".json",)),
        (("xml file", "xml"), (".xml",)),
    ]

    for phrases, extensions in type_rules:
        if any(phrase in question_lower for phrase in phrases):
            matches = [
                filename for filename in indexed_files
                if filename.lower().endswith(extensions)
            ]
            if len(matches) == 1:
                return matches[0]
            # If multiple files of that type exist, don't guess which one.
            return None

    return None


def generate_answer(prompt):
    """Generate an answer using the local Ollama model."""
    payload = json.dumps({
        "model": OLLAMA_MODEL,
        "prompt": prompt,
        "stream": False,
        "keep_alive": OLLAMA_KEEP_ALIVE,
    }).encode("utf-8")

    request = urllib.request.Request(
        f"{OLLAMA_URL}/api/generate",
        data=payload,
        headers={"Content-Type": "application/json"},
        method="POST",
    )

    with urllib.request.urlopen(request, timeout=180) as response:
        result = json.loads(response.read().decode("utf-8"))

    return result.get("response", "").strip()


db = get_db()

uploaded_files = st.file_uploader(
    "Upload your files",
    type=SUPPORTED_TYPES,
    accept_multiple_files=True,
    help="Upload one or more supported documents or images.",
)

if st.button("Process documents", type="primary"):
    if not uploaded_files:
        st.warning("Upload at least one file first.")
    else:
        all_records = []
        failed_files = []

        progress = st.progress(0)
        status = st.empty()

        for index, uploaded_file in enumerate(uploaded_files):
            status.write(
                f"Extracting {uploaded_file.name}..."
            )
            try:
                records = extract_document(
                    uploaded_file.getvalue(),
                    uploaded_file.name,
                )
                all_records.extend(records)

                if not records:
                    failed_files.append(
                        f"{uploaded_file.name}: no readable text found"
                    )
            except Exception as error:
                failed_files.append(
                    f"{uploaded_file.name}: {error}"
                )

            progress.progress((index + 1) / len(uploaded_files))

        chunks = split_records(all_records)

        if not chunks:
            st.error(
                "No readable text was extracted. "
                "Check the file contents and OCR language data."
            )
        else:
            try:
                status.write("Creating embeddings...")
                embedding_model = get_model()
                vectors = embedding_model.encode(
                    [chunk["text"] for chunk in chunks],
                    show_progress_bar=False,
                ).tolist()

                try:
                    db.delete_collection("documents")
                except Exception:
                    pass

                collection = db.create_collection("documents")
                collection.add(
                    ids=[str(i) for i in range(len(chunks))],
                    documents=[chunk["text"] for chunk in chunks],
                    embeddings=vectors,
                    metadatas=[
                        {
                            "source": chunk["source"],
                            "location": chunk["location"],
                        }
                        for chunk in chunks
                    ],
                )

                st.session_state["ready"] = True
                st.session_state["files"] = sorted({
                    chunk["source"] for chunk in chunks
                })

                st.success(
                    f"Indexed {len(chunks)} chunks from "
                    f"{len(st.session_state['files'])} file(s)."
                )

                if failed_files:
                    st.warning(
                        "Some files could not be fully processed:\n\n"
                        + "\n".join(failed_files)
                    )

                status.empty()
                progress.empty()

            except Exception as error:
                st.error(f"Indexing failed: {error}")


if st.session_state.get("ready"):
    st.caption(
        "Indexed files: "
        + ", ".join(st.session_state.get("files", []))
    )

    question = st.text_input(
        "Ask a question about your documents"
    )

    if st.button("Ask AI", type="primary") and question.strip():
        with st.spinner(
            "Searching your documents and generating an answer..."
        ):
            try:
                collection = db.get_collection("documents")
                count = collection.count()

                if count == 0:
                    st.warning("Process documents first.")
                    st.stop()

                embedding_model = get_model()
                query_vector = embedding_model.encode(
                    [question]
                ).tolist()

                indexed_files = st.session_state.get("files", [])
                requested_source = resolve_requested_source(
                    question, indexed_files
                )

                if requested_source:
                    source_filter = {"source": {"$eq": requested_source}}
                    source_records = collection.get(
                        where=source_filter,
                        include=["metadatas"],
                    )
                    source_count = len(source_records.get("ids", []))

                    if source_count == 0:
                        st.warning(
                            f"No indexed passages found for {requested_source}. "
                            "Please process that file again."
                        )
                        st.stop()

                    results = collection.query(
                        query_embeddings=query_vector,
                        n_results=min(8, source_count),
                        where=source_filter,
                    )
                    st.caption(
                        f"Searching only in: {requested_source}"
                    )
                else:
                    results = collection.query(
                        query_embeddings=query_vector,
                        n_results=min(4, count),
                    )

                passages = results["documents"][0]
                sources = results["metadatas"][0]

                evidence_parts = []
                for passage, source in zip(passages, sources):
                    evidence_parts.append(
                        f"[Source: {source['source']}; "
                        f"Location: {source['location']}]\n"
                        f"{passage}"
                    )

                evidence = "\n\n".join(evidence_parts)

                prompt = f"""
You are a careful assistant answering questions using retrieved document evidence.

Rules:
- Use only facts explicitly supported by the evidence below. Do not use outside knowledge.
- Keep each source separate. Never combine a person's name, budget, date, or other fact from one source with a different source's facts.
- If the question names a particular file, answer only from evidence belonging to that file.
- If the evidence does not contain every requested detail, say exactly which detail you could not find instead of guessing.
- Cite the filename and location for each factual answer.
- Before answering, check that the name and value you provide are supported by the same relevant source passage(s).

Evidence:
{evidence}

Question: {question}

Answer:
"""

                answer = generate_answer(prompt)

                st.subheader("Answer")
                st.write(
                    answer or "The model returned an empty answer."
                )

                st.subheader("Retrieved sources")

                for passage, source in zip(passages, sources):
                    label = (
                        f"{source['source']} — "
                        f"{source['location']}"
                    )
                    with st.expander(label):
                        st.write(passage)

            except urllib.error.HTTPError as error:
                details = error.read().decode(
                    "utf-8", errors="replace"
                )
                st.error(
                    f"Ollama returned HTTP {error.code}: {details}"
                )

            except urllib.error.URLError:
                st.error(
                    "Cannot connect to Ollama. Check that its service "
                    f"is running at {OLLAMA_URL}."
                )

            except Exception as error:
                st.error(f"Something went wrong: {error}")
